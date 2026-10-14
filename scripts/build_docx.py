import re, os
import pandas as pd
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
for sec in doc.sections:
    sec.left_margin = sec.right_margin = Inches(1)
style = doc.styles['Normal']; style.font.name='Times New Roman'; style.font.size=Pt(11)

def add_runs(p, text):
    # bold **...**, italics *...*, superscript [n] citation markers
    i=0
    for m in re.finditer(r'(\*\*.+?\*\*|\*[^*]+?\*|\[[0-9,\- ]+\])', text):
        if m.start()>i: p.add_run(text[i:m.start()])
        tok=m.group(0)
        if tok.startswith('**'): r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('['): r=p.add_run(tok); r.font.superscript=True
        else: r=p.add_run(tok[1:-1]); r.italic=True
        i=m.end()
    if i<len(text): p.add_run(text[i:])

def para(text, style_name=None):
    p=doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()
    add_runs(p, text); return p

def table_from_df(df, font_size=8):
    t=doc.add_table(rows=1+len(df), cols=len(df.columns)); t.style='Table Grid'
    for j,c in enumerate(df.columns):
        cell=t.rows[0].cells[j]; r=cell.paragraphs[0].add_run(str(c)); r.bold=True; r.font.size=Pt(font_size)
    for i,row in enumerate(df.itertuples(index=False),1):
        for j,v in enumerate(row):
            p=t.rows[i].cells[j].paragraphs[0]; r=p.add_run(str(v)); r.font.size=Pt(font_size)
    return t

lines = open('manuscript/manuscript.md').read().split('\n')
figures_inserted = set()
i=0
in_refs = False
md_table_rows = []
def flush_md_table():
    global md_table_rows
    rows=[r for r in md_table_rows if not set(r.replace('|','').strip())<=set('-: ')]
    if rows:
        cells=[[c.strip() for c in r.strip('|').split('|')] for r in rows]
        t=doc.add_table(rows=len(cells), cols=len(cells[0])); t.style='Table Grid'
        for ri,row in enumerate(cells):
            for ci,v in enumerate(row):
                rr=t.rows[ri].cells[ci].paragraphs[0].add_run(v); rr.font.size=Pt(8)
                if ri==0: rr.bold=True
    md_table_rows=[]
while i < len(lines):
    ln = lines[i]
    if ln.strip().startswith('|'):
        md_table_rows.append(ln); i+=1; continue
    if md_table_rows: flush_md_table()
    if ln.startswith('## '):
        in_refs = ln.strip().startswith('## References')
    if ln.startswith('# '):
        h=doc.add_paragraph(); r=h.add_run(ln[2:].strip()); r.bold=True; r.font.size=Pt(16)
        h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif ln.startswith('## '):
        h=doc.add_paragraph(); r=h.add_run(ln[3:].strip()); r.bold=True; r.font.size=Pt(13)
    elif ln.startswith('### '):
        h=doc.add_paragraph(); r=h.add_run(ln[4:].strip()); r.bold=True; r.font.size=Pt(11.5)
    elif ln.startswith('**'):
        para(ln)
    elif ln.strip()=='---':
        pass
    elif ln.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); add_runs(p, ln[2:])
    elif re.match(r'^\d+\. ', ln):
        if in_refs:
            para(ln)  # keep literal Vancouver numbers; Word auto-numbering offsets them
        else:
            p=doc.add_paragraph(style='List Number'); add_runs(p, re.sub(r'^\d+\. ','',ln))
    elif ln.strip():
        para(ln)
    i+=1

# Figures inline at end (journal style: figures as separate files too; inline for review copy)
doc.add_page_break()
h=doc.add_paragraph(); r=h.add_run('Figures'); r.bold=True; r.font.size=Pt(13)
captions={
 'fig1_pipeline.png':'Figure 1. Three-stage pipeline and staged conservative scan.',
 'fig2_family_map.png':'Figure 2. Cross-model primary interior-optimum landscape (mid-glycolysis family).',
 'fig2_efficiency_nopm.png':'Figure 3. Efficiency interior optima: ATP per glucose vs modulation intensity u, both models.',
 'fig3_pareto_gapdh.png':'Figure 4. Rate-yield trade-off: ATP rate vs ATP/glucose under GAPDH restriction.',
 'fig5_robustness.png':'Figure 5. Selector-by-environment robustness matrix (pFBA vs L1-MOMA).',
 'fig4_parsimony_nopm.png':'Figure S1. Secondary optima: interior minima in total |flux| (state-selection-consistent; not independent NOPM evidence).'}
for fn,cap in captions.items():
    fp=f'figures/{fn}'
    if os.path.exists(fp):
        doc.add_picture(fp, width=Inches(5.5))
        doc.paragraphs[-1].alignment=WD_ALIGN_PARAGRAPH.CENTER
        para(cap)

# Tables inline
doc.add_page_break()
h=doc.add_paragraph(); r=h.add_run('Main tables'); r.bold=True; r.font.size=Pt(13)
para('Table 1. Primary independent interior optima: model, reaction, family, endpoint, u*, J(0), J(u*), J(1), relative gain, cross-model family status.')
t1=pd.read_csv('manuscript/table1_primary_nopm.csv'); table_from_df(t1)
para('Table 2. Robustness matrix: interior-optimum counts per cell (model x selector x medium x modulation method), primary and secondary combined.')
t2=pd.read_csv('manuscript/table2_robustness.csv'); table_from_df(t2)
doc.add_page_break()
h=doc.add_paragraph(); r=h.add_run('Supplemental tables'); r.bold=True; r.font.size=Pt(13)
para('Table S1. Reaction crosswalk, Human-GEM <-> Recon3D.')
ts0=pd.read_csv('manuscript/table_crosswalk.csv'); table_from_df(ts0, 7)
para('Table S2. Prior-art comparison.')
ts2=pd.read_csv('manuscript/table3_prior_art.csv') if os.path.exists('manuscript/table3_prior_art.csv') else pd.DataFrame([
 ['Schäffner et al. 2026','Enzyme kinetics, single enzyme','No'],
 ['Motter 2010','Network damage/rescue','No'],
 ['Braess et al. 2005','Capacity addition paradox','No'],
 ['This work','Genome-scale interior-optimum census, human models','Yes']])
ts2.columns=list(ts2.columns) if len(list(ts2.columns))>0 else []
table_from_df(ts2)
para('Table S3. Qualifying targets and representative modulators (modulation\u2192target map, no dose translation; near-threshold nonqualifying targets in drug_mapping_nonqualifying.csv).')
ts3=pd.read_csv('data/drug_mapping.csv'); table_from_df(ts3, 7)


os.makedirs('submission',exist_ok=True)
doc.save('submission/manuscript_inline.docx')
print('saved submission/manuscript_inline.docx')
